from __future__ import annotations

import csv
import hashlib
import io
import json
import os
import re
import secrets
import socket
import ipaddress
from dataclasses import dataclass
from datetime import timedelta
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import urlparse

import httpx
from openpyxl import load_workbook

from platform_core import (
    ENVIRONMENT,
    IS_POSTGRES,
    MAX_UPLOAD_BYTES,
    PUBLIC_BASE_URL,
    audit,
    db,
    decrypt_secret_dict,
    encrypt_secret_dict,
    insert_id,
    logger,
    normalize,
    now_dt,
    now_iso,
    scalar,
    token_hash,
)

CANONICAL_FIELDS = (
    "brand",
    "part_number",
    "description",
    "condition",
    "quantity",
    "location",
)

FIELD_LABELS = {
    "brand": "Manufacturer / brand",
    "part_number": "Part number",
    "description": "Description",
    "condition": "Condition",
    "quantity": "Quantity",
    "location": "Location",
}

ALIASES = {
    "brand": {"brand", "manufacturer", "make", "mfr", "vendor", "maker"},
    "part_number": {"partnumber", "partno", "part", "sku", "productcode", "stockcode", "model", "itemcode", "material"},
    "description": {"description", "productdescription", "name", "productname", "details", "itemdescription"},
    "condition": {"condition", "state", "stockcondition", "quality", "grade"},
    "quantity": {"quantity", "qty", "stock", "onhand", "available", "stockqty", "freeqty", "freequantity"},
    "location": {"location", "warehouse", "site", "city", "stocklocation", "branch", "depot"},
}

SUPPORTED_TABULAR_EXTENSIONS = {".csv", ".xlsx", ".xlsm"}


@dataclass
class ParsedSource:
    headers: list[str]
    rows: list[dict[str, Any]]


def _header_key(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value or "").lower())


def read_tabular_file(filename: str, raw: bytes) -> ParsedSource:
    if len(raw) > MAX_UPLOAD_BYTES:
        raise ValueError(f"File is too large. Maximum upload size is {MAX_UPLOAD_BYTES // (1024 * 1024)} MB.")
    ext = Path(filename or "").suffix.lower()
    rows: list[dict[str, Any]] = []
    headers: list[str] = []
    if ext == ".csv":
        text = raw.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        headers = [str(h or "").strip() for h in (reader.fieldnames or [])]
        rows = [{str(k or "").strip(): v for k, v in dict(r).items()} for r in reader]
    elif ext in {".xlsx", ".xlsm"}:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        iterator = ws.iter_rows(values_only=True)
        try:
            headers = [str(v).strip() if v is not None else "" for v in next(iterator)]
        except StopIteration:
            return ParsedSource([], [])
        for values in iterator:
            rows.append({headers[i]: values[i] if i < len(values) else None for i in range(len(headers))})
    elif ext == ".xls":
        raise ValueError("Legacy .xls is not supported. Save/export it as .xlsx or .csv.")
    else:
        raise ValueError("Please use a CSV or XLSX file.")
    headers = [h for h in headers if h]
    if not headers:
        raise ValueError("The file has no usable header row.")
    return ParsedSource(headers, rows)


def guess_column_mapping(headers: Iterable[str]) -> dict[str, str]:
    result: dict[str, str] = {}
    for original in headers:
        key = _header_key(original)
        for canonical, aliases in ALIASES.items():
            if canonical not in result and key in aliases:
                result[canonical] = str(original)
                break
    return result


def mapping_from_form(form: Any) -> dict[str, str]:
    result = {}
    for field in CANONICAL_FIELDS:
        value = str(form.get(f"map_{field}", "") or "").strip()
        if value:
            result[field] = value
    return result


def _safe_json(value: Any) -> str:
    return json.dumps(value, default=str, separators=(",", ":"), ensure_ascii=False)


def transform_rows(rows: list[dict[str, Any]], mapping: dict[str, str]) -> tuple[list[dict[str, Any]], list[dict[str, Any]], list[dict[str, Any]]]:
    if not mapping.get("part_number"):
        raise ValueError("Map a column to Part number before importing.")
    cleaned: list[dict[str, Any]] = []
    errors: list[dict[str, Any]] = []
    warnings: list[dict[str, Any]] = []
    for index, row in enumerate(rows, start=2):
        part = str(row.get(mapping["part_number"], "") or "").strip()
        if not part:
            errors.append({"row_number": index, "code": "missing_part_number", "message": "Part number is blank.", "raw": row})
            continue
        q_header = mapping.get("quantity")
        qraw = row.get(q_header, 1) if q_header else 1
        if qraw in (None, ""):
            qty = 1
            warnings.append({"row_number": index, "code": "blank_quantity", "message": "Quantity was blank; defaulted to 1.", "raw": row})
        else:
            try:
                qty_float = float(str(qraw).replace(",", "").strip())
                if qty_float < 0 or int(qty_float) != qty_float:
                    raise ValueError
                qty = int(qty_float)
            except Exception:
                errors.append({"row_number": index, "code": "invalid_quantity", "message": f"Quantity '{qraw}' is not a non-negative whole number.", "raw": row})
                continue
        def value(field: str, default: str = "") -> str:
            header = mapping.get(field)
            return str(row.get(header, "") or default).strip() if header else default
        cleaned.append({
            "source_row": index,
            "brand": value("brand"),
            "part_number": part,
            "description": value("description"),
            "condition": value("condition", "Not stated") or "Not stated",
            "quantity": qty,
            "location": value("location"),
        })
    return cleaned, errors, warnings


def inventory_identity(item: dict[str, Any]) -> str:
    raw = "|".join([
        normalize(str(item.get("brand", ""))),
        normalize(str(item.get("part_number", ""))),
        normalize(str(item.get("condition", ""))),
        normalize(str(item.get("location", ""))),
    ])
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()[:40]


def preview_duplicates(items: list[dict[str, Any]], company_id: int, *, feed_id: int | None = None) -> dict[str, Any]:
    counts: dict[str, int] = {}
    for item in items:
        key = inventory_identity(item)
        counts[key] = counts.get(key, 0) + 1
    in_file = sum(v - 1 for v in counts.values() if v > 1)
    keys = list(counts)
    existing_keys: set[str] = set()
    with db() as conn:
        if feed_id:
            rows = conn.execute("SELECT source_key FROM inventory WHERE company_id=? AND source_feed_id=? AND deleted_at IS NULL", (company_id, feed_id)).fetchall()
        else:
            rows = conn.execute("SELECT source_key,brand,part_number,condition,location FROM inventory WHERE company_id=? AND source_feed_id IS NULL AND deleted_at IS NULL", (company_id,)).fetchall()
        for row in rows:
            key = row["source_key"] if row["source_key"] else inventory_identity(dict(row))
            if key:
                existing_keys.add(key)
    return {"within_file": in_file, "existing_matches": sum(1 for k in keys if k in existing_keys), "unique_rows": len(keys)}


def _record_issue(conn, job_id: int, issue: dict[str, Any], severity: str) -> None:
    conn.execute(
        "INSERT INTO inventory_import_errors(job_id,row_number,severity,code,message,raw_json,created_at) VALUES(?,?,?,?,?,?,?)",
        (job_id, issue.get("row_number"), severity, issue.get("code", ""), issue.get("message", ""), _safe_json(issue.get("raw", {})), now_iso()),
    )


def _merge_incoming(items: list[dict[str, Any]], policy: str) -> tuple[list[dict[str, Any]], int, list[dict[str, Any]]]:
    if policy == "keep":
        kept = []
        seen: dict[str, int] = {}
        for item in items:
            copy = dict(item)
            base = inventory_identity(copy)
            seen[base] = seen.get(base, 0) + 1
            copy["source_key"] = f"{base}:{seen[base]}"
            kept.append(copy)
        return kept, sum(v - 1 for v in seen.values() if v > 1), []
    by_key: dict[str, dict[str, Any]] = {}
    duplicate_count = 0
    issues = []
    for item in items:
        key = inventory_identity(item)
        if key not in by_key:
            copy = dict(item)
            copy["source_key"] = key
            by_key[key] = copy
            continue
        duplicate_count += 1
        issues.append({"row_number": item.get("source_row"), "code": "duplicate_in_file", "message": f"Duplicate of {item['part_number']} in the same import file.", "raw": item})
        if policy == "merge":
            by_key[key]["quantity"] += item["quantity"]
            if item.get("description") and not by_key[key].get("description"):
                by_key[key]["description"] = item["description"]
        # policy=skip keeps the first row unchanged
    return list(by_key.values()), duplicate_count, issues


def run_import(
    *,
    company_id: int,
    user_id: int | None,
    source_type: str,
    source_name: str,
    raw_rows: list[dict[str, Any]],
    mapping: dict[str, str],
    import_mode: str = "upsert",
    duplicate_policy: str = "merge",
    feed_id: int | None = None,
    default_location: str = "",
) -> tuple[int, list[int], dict[str, int]]:
    if import_mode not in {"add", "upsert", "sync", "replace"}:
        raise ValueError("Invalid import mode")
    if duplicate_policy not in {"merge", "skip", "keep"}:
        raise ValueError("Invalid duplicate policy")
    if import_mode == "sync" and feed_id and not raw_rows:
        raise ValueError("Feed returned zero rows. Sync was cancelled to protect the existing catalogue.")
    cleaned, errors, warnings = transform_rows(raw_rows, mapping)
    processed, duplicate_count, duplicate_issues = _merge_incoming(cleaned, duplicate_policy)
    warnings.extend(duplicate_issues)
    # Phase 5 commercial access and inventory-cap enforcement. Pending-verification suppliers
    # can prepare their catalogue before the free trial starts; expired accounts retain data
    # but cannot add/sync more stock until access is restored.
    from billing import catalogue_management_allowed, effective_plan
    with db() as limit_conn:
        if not catalogue_management_allowed(limit_conn, company_id):
            raise ValueError("Commercial access is paused. Choose a plan in Billing before importing inventory.")
        plan = effective_plan(limit_conn, company_id)
        inventory_limit = int(plan.get("inventory_limit", 0) or 0)
        if inventory_limit > 0:
            current_rows = limit_conn.execute("SELECT source_feed_id,source_key,brand,part_number,condition,location FROM inventory WHERE company_id=? AND active=1 AND deleted_at IS NULL", (company_id,)).fetchall()
            current_total = len(current_rows)
            incoming_keys = {item.get("source_key") or inventory_identity(item) for item in processed}
            if import_mode == "replace":
                projected = len(processed)
            else:
                if feed_id:
                    scope = [r for r in current_rows if r["source_feed_id"] == feed_id]
                else:
                    scope = [r for r in current_rows if r["source_feed_id"] is None]
                existing_keys = {r["source_key"] if r["source_key"] else inventory_identity(dict(r)) for r in scope}
                if import_mode == "sync" and feed_id and not errors:
                    projected = current_total - len(scope) + len(incoming_keys)
                else:
                    projected = current_total + sum(1 for key in incoming_keys if key not in existing_keys)
            if projected > inventory_limit:
                raise ValueError(f"This import would take the catalogue to about {projected:,} active lines, above the {inventory_limit:,}-line limit on {plan['name']}. Upgrade the plan or reduce the import.")
    started = now_iso()
    changed_ids: list[int] = []
    with db() as conn:
        job_id = insert_id(conn, """INSERT INTO inventory_import_jobs(company_id,user_id,feed_id,source_type,source_name,import_mode,duplicate_policy,status,total_rows,started_at)
            VALUES(?,?,?,?,?,?,?,?,?,?)""", (company_id, user_id, feed_id, source_type, source_name, import_mode, duplicate_policy, "running", len(raw_rows), started))
        for issue in errors:
            _record_issue(conn, job_id, issue, "error")
        for issue in warnings:
            _record_issue(conn, job_id, issue, "warning")

        now = now_iso()
        archived = inserted = updated = 0
        skipped = duplicate_count if duplicate_policy == "skip" else 0
        if import_mode == "replace":
            result_rows = conn.execute("SELECT id FROM inventory WHERE company_id=? AND active=1 AND deleted_at IS NULL", (company_id,)).fetchall()
            archived = len(result_rows)
            conn.execute("UPDATE inventory SET active=0,archived_at=?,updated_at=?,updated_by_user_id=? WHERE company_id=? AND active=1 AND deleted_at IS NULL", (now, now, user_id, company_id))

        incoming_keys: set[str] = set()
        for item in processed:
            key = item.get("source_key") or inventory_identity(item)
            incoming_keys.add(key)
            location = item.get("location") or default_location
            params: list[Any] = [company_id]
            sql = "SELECT * FROM inventory WHERE company_id=? AND deleted_at IS NULL"
            if feed_id:
                sql += " AND source_feed_id=?"
                params.append(feed_id)
            else:
                sql += " AND source_feed_id IS NULL"
            sql += " AND source_key=? ORDER BY active DESC,id ASC LIMIT 1"
            params.append(key)
            existing = None if import_mode in {"add", "replace"} else conn.execute(sql, tuple(params)).fetchone()
            if existing is None and import_mode not in {"add", "replace"}:
                # Phase 1/2 inventory did not have source_key metadata. Match legacy rows by
                # normalized part and full identity once, then stamp the key for future syncs.
                legacy_sql = "SELECT * FROM inventory WHERE company_id=? AND deleted_at IS NULL AND normalized_part=?"
                legacy_params: list[Any] = [company_id, normalize(item["part_number"])]
                if feed_id:
                    legacy_sql += " AND source_feed_id=?"; legacy_params.append(feed_id)
                else:
                    legacy_sql += " AND source_feed_id IS NULL"
                for candidate in conn.execute(legacy_sql + " ORDER BY active DESC,id ASC", tuple(legacy_params)).fetchall():
                    if inventory_identity(dict(candidate)) == key:
                        existing = candidate
                        break
            if existing:
                conn.execute("""UPDATE inventory SET brand=?,part_number=?,normalized_part=?,description=?,condition=?,quantity=?,location=?,active=1,
                    archived_at=NULL,last_confirmed_at=?,updated_at=?,updated_by_user_id=?,source_feed_id=?,source_key=?,last_import_job_id=? WHERE id=?""",
                    (item["brand"], item["part_number"], normalize(item["part_number"]), item["description"], item["condition"], item["quantity"], location,
                     now, now, user_id, feed_id, key, job_id, existing["id"]))
                changed_ids.append(existing["id"])
                updated += 1
            else:
                item_id = insert_id(conn, """INSERT INTO inventory(company_id,brand,part_number,normalized_part,description,condition,quantity,location,active,
                    created_at,created_by_user_id,updated_by_user_id,last_confirmed_at,updated_at,source_feed_id,source_key,last_import_job_id)
                    VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                    (company_id, item["brand"], item["part_number"], normalize(item["part_number"]), item["description"], item["condition"], item["quantity"], location,
                     1, now, user_id, user_id, now, now, feed_id, key, job_id))
                changed_ids.append(item_id)
                inserted += 1

        retirement_suppressed = False
        if import_mode == "sync" and feed_id:
            if errors:
                retirement_suppressed = True
                warning = {"row_number": None, "code": "sync_retirement_suppressed", "message": "Some source rows were invalid, so missing feed lines were not retired during this run.", "raw": {}}
                warnings.append(warning)
                _record_issue(conn, job_id, warning, "warning")
            else:
                feed_rows = conn.execute("SELECT id,source_key FROM inventory WHERE company_id=? AND source_feed_id=? AND active=1 AND deleted_at IS NULL", (company_id, feed_id)).fetchall()
                stale_ids = [r["id"] for r in feed_rows if r["source_key"] not in incoming_keys]
                for item_id in stale_ids:
                    conn.execute("UPDATE inventory SET active=0,archived_at=?,updated_at=?,updated_by_user_id=?,last_import_job_id=? WHERE id=?", (now, now, user_id, job_id, item_id))
                archived += len(stale_ids)

        status = "completed_with_errors" if errors else "completed"
        report = {
            "mapping": mapping,
            "warnings": len(warnings),
            "errors": len(errors),
            "changed_inventory_ids": changed_ids[:5000],
            "sync_retirement_suppressed": retirement_suppressed,
        }
        conn.execute("""UPDATE inventory_import_jobs SET status=?,inserted_rows=?,updated_rows=?,archived_rows=?,duplicate_rows=?,skipped_rows=?,error_rows=?,
            report_json=?,error_summary=?,completed_at=? WHERE id=?""",
            (status, inserted, updated, archived, duplicate_count, skipped, len(errors), _safe_json(report), errors[0]["message"] if errors else "", now_iso(), job_id))
        audit(conn, None, {"id": user_id, "company_id": company_id} if user_id else None, "inventory.imported", "inventory_import_job", job_id,
              {"source_type": source_type, "source_name": source_name, "inserted": inserted, "updated": updated, "archived": archived, "errors": len(errors)}, company_id=company_id)
    try:
        from catalog_intelligence import link_inventory_ids
        link_inventory_ids(changed_ids)
    except Exception:
        logger.exception("catalog_link_after_import_failed job_id=%s", job_id)
    return job_id, changed_ids, {"inserted": inserted, "updated": updated, "archived": archived, "duplicates": duplicate_count, "skipped": skipped, "errors": len(errors), "warnings": len(warnings)}


def run_file_import(**kwargs):
    filename = kwargs.pop("filename")
    raw = kwargs.pop("raw")
    mapping = kwargs.pop("mapping", None)
    parsed = read_tabular_file(filename, raw)
    mapping = mapping or guess_column_mapping(parsed.headers)
    if not mapping.get("part_number"):
        raise ValueError("Could not identify a part-number column. Use the mapping screen to choose it manually.")
    return run_import(raw_rows=parsed.rows, mapping=mapping, source_name=filename, **kwargs)


def create_staging(company_id: int, user_id: int, filename: str, raw: bytes) -> tuple[int, ParsedSource, dict[str, str]]:
    parsed = read_tabular_file(filename, raw)
    guessed = guess_column_mapping(parsed.headers)
    expires = (now_dt() + timedelta(hours=2)).isoformat(timespec="seconds")
    with db() as conn:
        conn.execute("DELETE FROM inventory_import_staging WHERE expires_at<?", (now_iso(),))
        stage_id = insert_id(conn, "INSERT INTO inventory_import_staging(company_id,user_id,filename,file_data,headers_json,expires_at,created_at) VALUES(?,?,?,?,?,?,?)",
            (company_id, user_id, filename, raw, _safe_json(parsed.headers), expires, now_iso()))
    return stage_id, parsed, guessed


def load_staging(stage_id: int, company_id: int, user_id: int) -> tuple[Any, ParsedSource]:
    with db() as conn:
        stage = conn.execute("SELECT * FROM inventory_import_staging WHERE id=? AND company_id=? AND user_id=? AND expires_at>=?", (stage_id, company_id, user_id, now_iso())).fetchone()
    if not stage:
        raise ValueError("That import preview expired. Upload the file again.")
    raw = bytes(stage["file_data"])
    parsed = read_tabular_file(stage["filename"], raw)
    return stage, parsed


def delete_staging(stage_id: int) -> None:
    with db() as conn:
        conn.execute("DELETE FROM inventory_import_staging WHERE id=?", (stage_id,))


def make_ingest_token() -> str:
    return secrets.token_urlsafe(32)


def feed_public_config(feed: Any) -> dict[str, Any]:
    try:
        return json.loads(feed["config_json"] or "{}")
    except Exception:
        return {}


def feed_secrets(feed: Any) -> dict[str, Any]:
    return decrypt_secret_dict(feed["secret_blob"])


def _http_headers(secrets_data: dict[str, Any]) -> dict[str, str]:
    headers = {}
    bearer = str(secrets_data.get("bearer_token", "") or "").strip()
    if bearer:
        headers["Authorization"] = f"Bearer {bearer}"
    header_name = str(secrets_data.get("header_name", "") or "").strip()
    header_value = str(secrets_data.get("header_value", "") or "").strip()
    if header_name and header_value:
        headers[header_name] = header_value
    return headers


def _safe_remote_host(hostname: str) -> None:
    if ENVIRONMENT != "production" or os.getenv("ALLOW_PRIVATE_FEED_HOSTS", "false").lower() == "true":
        return
    try:
        infos = socket.getaddrinfo(hostname, None, type=socket.SOCK_STREAM)
    except socket.gaierror as exc:
        raise ValueError(f"Could not resolve remote feed host: {hostname}") from exc
    if not infos:
        raise ValueError("Remote feed host did not resolve.")
    for info in infos:
        address = info[4][0].split("%", 1)[0]
        ip = ipaddress.ip_address(address)
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved or ip.is_multicast or ip.is_unspecified:
            raise ValueError("Remote feed host resolves to a private or non-public network address. Set ALLOW_PRIVATE_FEED_HOSTS=true only when intentionally using a protected private network.")


def _safe_remote_url(url: str) -> None:
    parsed = urlparse(url)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname:
        raise ValueError("Feed URL must be an http(s) URL.")
    if parsed.username or parsed.password:
        raise ValueError("Do not put credentials in the feed URL; use the encrypted authentication fields.")
    if ENVIRONMENT == "production" and parsed.scheme != "https":
        raise ValueError("Production remote feeds must use HTTPS.")
    _safe_remote_host(parsed.hostname)


def fetch_feed_payload(feed: Any) -> tuple[str, bytes | None, list[dict[str, Any]] | None]:
    feed_type = feed["feed_type"]
    config = feed_public_config(feed)
    secrets_data = feed_secrets(feed)
    if feed_type in {"url", "api_pull"}:
        url = str(config.get("url", "") or "").strip()
        _safe_remote_url(url)
        headers = _http_headers(secrets_data)
        forbidden = {"host", "content-length", "connection", "transfer-encoding"}
        if any(name.lower() in forbidden for name in headers):
            raise ValueError("That custom HTTP header is not allowed.")
        response = httpx.get(url, headers=headers, timeout=30, follow_redirects=False)
        if 300 <= response.status_code < 400:
            raise ValueError("Remote feed redirects are disabled for security. Configure the final HTTPS URL directly.")
        response.raise_for_status()
        if len(response.content) > MAX_UPLOAD_BYTES:
            raise ValueError("Remote feed exceeds the maximum import size.")
        if feed_type == "api_pull":
            payload = response.json()
            data_path = str(config.get("data_path", "") or "").strip()
            if data_path:
                for part in data_path.split("."):
                    payload = payload[part]
            if isinstance(payload, dict) and isinstance(payload.get("items"), list):
                payload = payload["items"]
            if not isinstance(payload, list) or not all(isinstance(x, dict) for x in payload):
                raise ValueError("API response must be a JSON array of objects (or {items:[...]}).")
            return str(config.get("filename_hint") or "api.json"), None, payload
        name = str(config.get("filename_hint") or Path(urlparse(url).path).name or "feed.csv")
        return name, response.content, None
    if feed_type == "sftp":
        try:
            import paramiko
        except ImportError as exc:
            raise RuntimeError("SFTP support requires the paramiko package from requirements.txt.") from exc
        host = str(config.get("host", "") or "").strip()
        username = str(config.get("username", "") or "").strip()
        remote_path = str(config.get("remote_path", "") or "").strip()
        port = int(config.get("port") or 22)
        if not host or not username or not remote_path:
            raise ValueError("SFTP host, username and remote path are required.")
        _safe_remote_host(host)
        expected_fp = str(config.get("host_key_sha256", "") or "").strip().replace("SHA256:", "")
        client = paramiko.SSHClient()
        client.load_system_host_keys()
        if expected_fp:
            class FingerprintPolicy(paramiko.MissingHostKeyPolicy):
                def missing_host_key(self, client_, hostname, key):
                    import base64
                    fp = base64.b64encode(hashlib.sha256(key.asbytes()).digest()).decode().rstrip("=")
                    if fp != expected_fp:
                        raise paramiko.SSHException("SFTP host key fingerprint does not match")
                    client_._host_keys.add(hostname, key.get_name(), key)
            client.set_missing_host_key_policy(FingerprintPolicy())
        else:
            if ENVIRONMENT == "production":
                raise ValueError("SFTP host_key_sha256 is required in production.")
            client.set_missing_host_key_policy(paramiko.AutoAddPolicy())
        connect_kwargs: dict[str, Any] = {"hostname": host, "port": port, "username": username, "timeout": 20, "banner_timeout": 20, "auth_timeout": 20}
        if secrets_data.get("password"):
            connect_kwargs["password"] = secrets_data["password"]
        if secrets_data.get("private_key"):
            key_text = str(secrets_data["private_key"])
            key_file = io.StringIO(key_text)
            parsed_key = None
            for cls in (paramiko.Ed25519Key, paramiko.RSAKey, paramiko.ECDSAKey):
                try:
                    key_file.seek(0)
                    parsed_key = cls.from_private_key(key_file, password=secrets_data.get("private_key_passphrase") or None)
                    break
                except Exception:
                    continue
            if not parsed_key:
                raise ValueError("Could not read the supplied SFTP private key.")
            connect_kwargs["pkey"] = parsed_key
        client.connect(**connect_kwargs)
        try:
            with client.open_sftp() as sftp:
                with sftp.open(remote_path, "rb") as remote:
                    raw = remote.read(MAX_UPLOAD_BYTES + 1)
            if len(raw) > MAX_UPLOAD_BYTES:
                raise ValueError("SFTP feed exceeds the maximum import size.")
        finally:
            client.close()
        return str(config.get("filename_hint") or Path(remote_path).name or "feed.csv"), raw, None
    raise ValueError(f"Feed type '{feed_type}' is push-only and cannot be polled.")


def run_feed(feed_id: int, *, actor_user_id: int | None = None) -> tuple[int, list[int], dict[str, int]]:
    with db() as conn:
        feed = conn.execute("SELECT f.*,c.location AS company_location FROM inventory_feeds f JOIN companies c ON c.id=f.company_id WHERE f.id=?", (feed_id,)).fetchone()
    if not feed:
        raise ValueError("Feed not found")
    filename, raw, json_rows = fetch_feed_payload(feed)
    mapping = json.loads(feed["mapping_json"] or "{}")
    if json_rows is not None:
        headers = list(dict.fromkeys(str(k) for row in json_rows[:100] for k in row.keys()))
        mapping = mapping or guess_column_mapping(headers)
        result = run_import(company_id=feed["company_id"], user_id=actor_user_id, feed_id=feed_id, source_type=feed["feed_type"], source_name=filename,
            raw_rows=json_rows, mapping=mapping, import_mode=feed["import_mode"], duplicate_policy=feed["duplicate_policy"], default_location=feed["company_location"] or "")
    else:
        if feed["import_mode"] == "sync" and not (raw or b""):
            raise ValueError("Feed returned zero rows. Sync was cancelled to protect the existing catalogue.")
        parsed = read_tabular_file(filename, raw or b"")
        mapping = mapping or guess_column_mapping(parsed.headers)
        result = run_import(company_id=feed["company_id"], user_id=actor_user_id, feed_id=feed_id, source_type=feed["feed_type"], source_name=filename,
            raw_rows=parsed.rows, mapping=mapping, import_mode=feed["import_mode"], duplicate_policy=feed["duplicate_policy"], default_location=feed["company_location"] or "")
    next_run = (now_dt() + timedelta(minutes=max(15, int(feed["schedule_minutes"] or 360)))).isoformat(timespec="seconds")
    with db() as conn:
        conn.execute("UPDATE inventory_feeds SET last_run_at=?,last_status='success',next_run_at=?,locked_at=NULL,updated_at=? WHERE id=?", (now_iso(), next_run, now_iso(), feed_id))
    return result


def mark_feed_failed(feed_id: int, message: str) -> None:
    with db() as conn:
        feed = conn.execute("SELECT schedule_minutes FROM inventory_feeds WHERE id=?", (feed_id,)).fetchone()
        if not feed:
            return
        retry_minutes = min(max(15, int(feed["schedule_minutes"] or 360)), 60)
        next_run = (now_dt() + timedelta(minutes=retry_minutes)).isoformat(timespec="seconds")
        conn.execute("UPDATE inventory_feeds SET last_run_at=?,last_status=?,next_run_at=?,locked_at=NULL,updated_at=? WHERE id=?", (now_iso(), ("error: " + message)[:500], next_run, now_iso(), feed_id))


def due_feed_ids(limit: int = 10) -> list[int]:
    with db() as conn:
        params = (now_iso(), (now_dt()-timedelta(minutes=20)).isoformat(timespec="seconds"), limit)
        if IS_POSTGRES:
            rows = conn.execute("""SELECT f.id FROM inventory_feeds f JOIN billing_subscriptions bs ON bs.company_id=f.company_id
                WHERE f.enabled=1 AND f.feed_type IN ('url','api_pull','sftp')
                AND (bs.status='pending_verification' OR bs.status IN ('active','past_due') OR (bs.status='trialing' AND bs.trial_ends_at>?))
                AND (f.next_run_at IS NULL OR f.next_run_at<=?) AND (f.locked_at IS NULL OR f.locked_at<?)
                ORDER BY COALESCE(f.next_run_at,f.created_at) ASC LIMIT ? FOR UPDATE SKIP LOCKED""", (now_iso(),) + params).fetchall()
        else:
            rows = conn.execute("""SELECT f.id FROM inventory_feeds f JOIN billing_subscriptions bs ON bs.company_id=f.company_id
                WHERE f.enabled=1 AND f.feed_type IN ('url','api_pull','sftp')
                AND (bs.status='pending_verification' OR bs.status IN ('active','past_due') OR (bs.status='trialing' AND bs.trial_ends_at>?))
                AND (f.next_run_at IS NULL OR f.next_run_at<=?) AND (f.locked_at IS NULL OR f.locked_at<?)
                ORDER BY COALESCE(f.next_run_at,f.created_at) ASC LIMIT ?""", (now_iso(),) + params).fetchall()
        ids = [r["id"] for r in rows]
        for feed_id in ids:
            conn.execute("UPDATE inventory_feeds SET locked_at=? WHERE id=?", (now_iso(), feed_id))
    return ids


def inbound_feed_by_token(raw_token: str, feed_types: tuple[str, ...]) -> Any | None:
    if not raw_token:
        return None
    placeholders = ",".join("?" for _ in feed_types)
    with db() as conn:
        return conn.execute(f"SELECT f.*,c.location AS company_location FROM inventory_feeds f JOIN companies c ON c.id=f.company_id WHERE f.enabled=1 AND f.feed_type IN ({placeholders}) AND f.ingest_token_hash=?", (*feed_types, token_hash(raw_token))).fetchone()


def inbound_address(feed: Any) -> str:
    from platform_core import IMPORT_EMAIL_DOMAIN
    if feed["feed_type"] != "email" or not IMPORT_EMAIL_DOMAIN:
        return ""
    secrets_data = feed_secrets(feed)
    token = secrets_data.get("inbound_token", "")
    return f"imports+{token}@{IMPORT_EMAIL_DOMAIN}" if token else ""
