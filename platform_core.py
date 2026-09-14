from __future__ import annotations

import base64
import csv
import difflib
import hashlib
import hmac
import io
import json
import logging
import os
import re
import secrets
import smtplib
import sqlite3
import ssl
from contextlib import contextmanager
from datetime import datetime, timedelta, timezone
from email.message import EmailMessage
from pathlib import Path
from typing import Any, Iterable, Optional

import httpx
from cryptography.fernet import Fernet, InvalidToken
from openpyxl import load_workbook

BASE_DIR = Path(__file__).resolve().parent
DATA_DIR = BASE_DIR / "data"
DATA_DIR.mkdir(parents=True, exist_ok=True)
BACKUP_DIR = Path(os.getenv("BACKUP_DIR", str(BASE_DIR / "backups")))
if not BACKUP_DIR.is_absolute():
    BACKUP_DIR = BASE_DIR / BACKUP_DIR
BACKUP_DIR.mkdir(parents=True, exist_ok=True)

ENVIRONMENT = os.getenv("ENVIRONMENT", "development").lower()
SECRET_KEY = os.getenv("SECRET_KEY", "dev-only-change-this-secret-key")
PUBLIC_BASE_URL = os.getenv("PUBLIC_BASE_URL", "http://127.0.0.1:8000").rstrip("/")
DATABASE_URL = os.getenv("DATABASE_URL", "").strip()
DATABASE_PATH = Path(os.getenv("DATABASE_PATH", str(DATA_DIR / "controls_exchange.db")))
if not DATABASE_PATH.is_absolute():
    DATABASE_PATH = BASE_DIR / DATABASE_PATH
DATABASE_PATH.parent.mkdir(parents=True, exist_ok=True)
IS_POSTGRES = DATABASE_URL.startswith("postgresql://") or DATABASE_URL.startswith("postgres://")
MAX_UPLOAD_BYTES = int(os.getenv("MAX_UPLOAD_BYTES", str(10 * 1024 * 1024)))
FRESH_DAYS = int(os.getenv("INVENTORY_FRESH_DAYS", "30"))
AGING_DAYS = int(os.getenv("INVENTORY_AGING_DAYS", "60"))
EXPIRY_DAYS = int(os.getenv("INVENTORY_EXPIRY_DAYS", "90"))
INTEGRATION_ENCRYPTION_KEY = os.getenv("INTEGRATION_ENCRYPTION_KEY", SECRET_KEY).strip()
IMPORT_EMAIL_DOMAIN = os.getenv("INVENTORY_IMPORT_EMAIL_DOMAIN", "").strip().lower()

COMPANY_ROLES = ("owner", "admin", "sales", "inventory", "viewer")
ROLE_PERMISSIONS = {
    "owner": {"manage_team", "manage_inventory", "trade", "view_audit"},
    "admin": {"manage_team", "manage_inventory", "trade", "view_audit"},
    "sales": {"trade"},
    "inventory": {"manage_inventory"},
    "viewer": set(),
}

logger = logging.getLogger("controls_exchange")


def validate_production_config() -> None:
    if ENVIRONMENT != "production":
        return
    problems = []
    if SECRET_KEY in {"dev-only-change-this-secret-key", "change-this-before-public-use"} or len(SECRET_KEY) < 32:
        problems.append("SECRET_KEY must be a random value of at least 32 characters")
    if not IS_POSTGRES:
        problems.append("DATABASE_URL must point to PostgreSQL in production")
    provider = os.getenv("EMAIL_PROVIDER", "").lower()
    if provider not in {"resend", "smtp"}:
        problems.append("EMAIL_PROVIDER must be 'resend' or 'smtp' in production")
    if not PUBLIC_BASE_URL.startswith("https://"):
        problems.append("PUBLIC_BASE_URL must use https:// in production")
    if len(os.getenv("INTEGRATION_ENCRYPTION_KEY", "")) < 32:
        problems.append("INTEGRATION_ENCRYPTION_KEY must be a random value of at least 32 characters in production")
    if not os.getenv("ORDER_DOCUMENT_DIR", "").strip():
        problems.append("ORDER_DOCUMENT_DIR must be set to persistent storage in production")
    try:
        doc_limit = int(os.getenv("ORDER_DOCUMENT_MAX_BYTES", str(10 * 1024 * 1024)))
        if doc_limit < 1024 * 1024 or doc_limit > 50 * 1024 * 1024:
            problems.append("ORDER_DOCUMENT_MAX_BYTES must be between 1 MB and 50 MB in production")
    except ValueError:
        problems.append("ORDER_DOCUMENT_MAX_BYTES must be an integer")
    stripe_secret = os.getenv("STRIPE_SECRET_KEY", "").strip()
    if stripe_secret and not os.getenv("STRIPE_WEBHOOK_SECRET", "").strip():
        problems.append("STRIPE_WEBHOOK_SECRET is required when STRIPE_SECRET_KEY is configured")
    if os.getenv("STRIPE_TEST_CLOCK_ID", "").strip():
        problems.append("STRIPE_TEST_CLOCK_ID is test-only and must not be set in production")
    if stripe_secret:
        required_prices = [
            "STRIPE_PRICE_SUPPLIER_STARTER_MONTHLY", "STRIPE_PRICE_SUPPLIER_STARTER_ANNUAL",
            "STRIPE_PRICE_SUPPLIER_PRO_MONTHLY", "STRIPE_PRICE_SUPPLIER_PRO_ANNUAL",
            "STRIPE_PRICE_SUPPLIER_PREMIUM_MONTHLY", "STRIPE_PRICE_SUPPLIER_PREMIUM_ANNUAL",
            "STRIPE_PRICE_MANUFACTURER_MONTHLY", "STRIPE_PRICE_MANUFACTURER_ANNUAL",
        ]
        missing_prices = [name for name in required_prices if not os.getenv(name, "").strip()]
        if missing_prices:
            problems.append("Stripe is enabled but price IDs are missing: " + ", ".join(missing_prices))
        portal_return = os.getenv("STRIPE_PORTAL_RETURN_URL", f"{PUBLIC_BASE_URL}/billing").strip()
        if not portal_return.startswith("https://"):
            problems.append("STRIPE_PORTAL_RETURN_URL must use https:// in production")
        if os.getenv("STRIPE_TRIAL_END_BEHAVIOR", "pause").strip().lower() not in {"pause", "cancel", "create_invoice"}:
            problems.append("STRIPE_TRIAL_END_BEHAVIOR must be pause, cancel, or create_invoice")
    # Phase 6 commercial intelligence must retain minimum aggregation cohorts in production.
    try:
        if int(os.getenv("INTELLIGENCE_MIN_DEMAND_EVENTS", "5")) < 5:
            problems.append("INTELLIGENCE_MIN_DEMAND_EVENTS must be at least 5 in production")
        if int(os.getenv("INTELLIGENCE_MIN_BUYER_COMPANIES", "3")) < 3:
            problems.append("INTELLIGENCE_MIN_BUYER_COMPANIES must be at least 3 in production")
        if int(os.getenv("INTELLIGENCE_MIN_QUOTE_SAMPLES", "5")) < 5:
            problems.append("INTELLIGENCE_MIN_QUOTE_SAMPLES must be at least 5 in production")
        if int(os.getenv("INTELLIGENCE_MIN_QUOTE_SUPPLIERS", "3")) < 3:
            problems.append("INTELLIGENCE_MIN_QUOTE_SUPPLIERS must be at least 3 in production")
        if int(os.getenv("INTELLIGENCE_MIN_QUOTE_BUYER_COMPANIES", "3")) < 3:
            problems.append("INTELLIGENCE_MIN_QUOTE_BUYER_COMPANIES must be at least 3 in production")
        share = float(os.getenv("INTELLIGENCE_MAX_CONTRIBUTOR_SHARE", "0.5"))
        if share > 0.5 or share <= 0:
            problems.append("INTELLIGENCE_MAX_CONTRIBUTOR_SHARE must be > 0 and <= 0.5 in production")
    except ValueError:
        problems.append("Phase 6 intelligence thresholds must be integers")
    if problems:
        raise RuntimeError("Unsafe production configuration: " + "; ".join(problems))



def _fernet() -> Fernet:
    digest = hashlib.sha256(INTEGRATION_ENCRYPTION_KEY.encode("utf-8")).digest()
    return Fernet(base64.urlsafe_b64encode(digest))


def encrypt_secret_dict(value: dict[str, Any] | None) -> str:
    if not value:
        return ""
    payload = json.dumps(value, separators=(",", ":"), sort_keys=True).encode("utf-8")
    return _fernet().encrypt(payload).decode("ascii")


def decrypt_secret_dict(value: str | None) -> dict[str, Any]:
    if not value:
        return {}
    try:
        raw = _fernet().decrypt(str(value).encode("ascii"))
        data = json.loads(raw.decode("utf-8"))
        return data if isinstance(data, dict) else {}
    except (InvalidToken, ValueError, json.JSONDecodeError):
        logger.warning("integration_secret_decrypt_failed")
        return {}


def generate_ingest_token() -> str:
    return secrets.token_urlsafe(32)


def now_dt() -> datetime:
    return datetime.now(timezone.utc)


def now_iso() -> str:
    return now_dt().isoformat(timespec="seconds")


def normalize(text: str | None) -> str:
    return re.sub(r"[^a-z0-9]", "", (text or "").lower())


def _row_value(row: Any, key: str, default: Any = "") -> Any:
    if row is None:
        return default
    try:
        return row[key]
    except Exception:
        return default


def parse_iso(value: str | None) -> datetime | None:
    if not value:
        return None
    try:
        parsed = datetime.fromisoformat(str(value).replace("Z", "+00:00"))
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=timezone.utc)
        return parsed.astimezone(timezone.utc)
    except Exception:
        return None


def freshness_details(value: str | None) -> dict[str, Any]:
    confirmed = parse_iso(value)
    if not confirmed:
        return {"state": "expired", "label": "Needs confirmation", "age_days": None, "searchable": False}
    age_days = max(0, int((now_dt() - confirmed).total_seconds() // 86400))
    if age_days <= FRESH_DAYS:
        state, label = "fresh", "Fresh"
    elif age_days <= AGING_DAYS:
        state, label = "aging", "Aging"
    elif age_days <= EXPIRY_DAYS:
        state, label = "stale", "Stale"
    else:
        state, label = "expired", "Expired"
    return {"state": state, "label": label, "age_days": age_days, "searchable": age_days <= EXPIRY_DAYS}


def freshness_cutoff_iso(days: int | None = None) -> str:
    days = EXPIRY_DAYS if days is None else days
    return (now_dt() - timedelta(days=days)).isoformat(timespec="seconds")


def part_segments(text: str | None) -> list[str]:
    return re.findall(r"[a-z]+|\d+", (text or "").lower())


def inventory_match_score(query: str, item: Any) -> tuple[float, str]:
    q = (query or "").strip().lower()
    qn = normalize(q)
    if not qn:
        return 0.0, ""
    part = str(_row_value(item, "part_number", "") or "")
    brand = str(_row_value(item, "brand", "") or "")
    desc = str(_row_value(item, "description", "") or "")
    pn = normalize(part)
    bn = normalize(brand)
    combined = f"{brand} {part} {desc}".lower()
    score = 0.0
    reason = "Related text"
    if qn == pn:
        return 100.0, "Exact part number"
    if pn and (pn.startswith(qn) or qn.startswith(pn)) and min(len(pn), len(qn)) >= 4:
        score, reason = 93.0, "Part-number prefix"
    if pn and qn in pn and len(qn) >= 3 and score < 90:
        score, reason = 88.0, "Part-number contains match"
    qseg = part_segments(q)
    pseg = part_segments(part)
    if qseg and all(seg in pseg for seg in qseg) and len(qseg) >= 2 and score < 86:
        score, reason = 86.0, "Part-number segments"
    if len(qn) >= 4 and pn:
        ratio = difflib.SequenceMatcher(None, qn, pn).ratio()
        if ratio >= 0.74:
            fuzzy = 58.0 + ratio * 30.0
            if fuzzy > score:
                score, reason = fuzzy, "Close part-number match"
    if qn == bn and score < 72:
        score, reason = 72.0, "Manufacturer match"
    elif q in brand.lower() and len(q) >= 3 and score < 68:
        score, reason = 68.0, "Manufacturer match"
    tokens = [t for t in re.findall(r"[a-z0-9]+", q) if len(t) > 1]
    if tokens:
        hits = sum(1 for t in tokens if t in combined)
        token_score = 45.0 + (hits / len(tokens)) * 22.0
        if hits and token_score > score:
            score, reason = token_score, "Description / manufacturer match"
    return round(score, 2), reason


def supplier_trust_metrics(conn: "Database", company_id: int) -> dict[str, Any]:
    stock = conn.execute("SELECT last_confirmed_at FROM inventory WHERE company_id=? AND active=1 AND deleted_at IS NULL AND quantity>0", (company_id,)).fetchall()
    active_stock = len(stock)
    fresh_stock = sum(1 for row in stock if freshness_details(_row_value(row, "last_confirmed_at"))["state"] == "fresh")
    searchable_stock = sum(1 for row in stock if freshness_details(_row_value(row, "last_confirmed_at"))["searchable"])
    latest = None
    for row in stock:
        dt = parse_iso(_row_value(row, "last_confirmed_at"))
        if dt and (latest is None or dt > latest):
            latest = dt
    rfq_rows = conn.execute("""SELECT rr.status, rr.responded_at, r.created_at
        FROM rfq_recipients rr JOIN rfqs r ON r.id=rr.rfq_id
        WHERE rr.supplier_company_id=?""", (company_id,)).fetchall()
    total_rfqs = len(rfq_rows)
    responded = [r for r in rfq_rows if _row_value(r, "responded_at") or _row_value(r, "status") in {"quoted","declined","accepted"}]
    response_hours = []
    for row in responded:
        a, b = parse_iso(_row_value(row, "created_at")), parse_iso(_row_value(row, "responded_at"))
        if a and b and b >= a:
            response_hours.append((b-a).total_seconds()/3600)
    accepted = sum(1 for r in rfq_rows if _row_value(r, "status") == "accepted")
    return {
        "active_stock": active_stock,
        "searchable_stock": searchable_stock,
        "fresh_stock_pct": round((fresh_stock / active_stock) * 100) if active_stock else 0,
        "rfq_response_rate": round((len(responded) / total_rfqs) * 100) if total_rfqs else None,
        "avg_response_hours": round(sum(response_hours)/len(response_hours), 1) if response_hours else None,
        "selected_quotes": accepted,
        "latest_inventory_at": latest.isoformat(timespec="seconds") if latest else None,
    }


def hash_password(password: str) -> str:
    salt = secrets.token_bytes(16)
    iterations = 310_000
    digest = hashlib.pbkdf2_hmac("sha256", password.encode(), salt, iterations)
    return f"pbkdf2_sha256${iterations}${salt.hex()}${digest.hex()}"


def verify_password(password: str, encoded: str) -> bool:
    try:
        algo, iterations, salt_hex, digest_hex = encoded.split("$", 3)
        if algo != "pbkdf2_sha256":
            return False
        digest = hashlib.pbkdf2_hmac(
            "sha256", password.encode(), bytes.fromhex(salt_hex), int(iterations)
        ).hex()
        return hmac.compare_digest(digest, digest_hex)
    except Exception:
        return False


def token_hash(raw: str) -> str:
    return hashlib.sha256(raw.encode()).hexdigest()


class DBResult:
    def __init__(self, cursor: Any, is_postgres: bool):
        self.cursor = cursor
        self.is_postgres = is_postgres
        self.lastrowid = None if is_postgres else cursor.lastrowid

    def fetchone(self):
        return self.cursor.fetchone()

    def fetchall(self):
        return self.cursor.fetchall()


class Database:
    def __init__(self):
        self.conn = None

    def __enter__(self):
        if IS_POSTGRES:
            try:
                import psycopg
                from psycopg.rows import dict_row
            except ImportError as exc:
                raise RuntimeError("PostgreSQL configured but psycopg is not installed. Run pip install -r requirements.txt") from exc
            url = DATABASE_URL.replace("postgres://", "postgresql://", 1)
            self.conn = psycopg.connect(url, row_factory=dict_row)
        else:
            self.conn = sqlite3.connect(DATABASE_PATH, timeout=30)
            self.conn.row_factory = sqlite3.Row
            self.conn.execute("PRAGMA foreign_keys = ON")
            self.conn.execute("PRAGMA journal_mode = WAL")
            self.conn.execute("PRAGMA busy_timeout = 30000")
        return self

    def __exit__(self, exc_type, exc, tb):
        if self.conn is None:
            return False
        try:
            if exc_type:
                self.conn.rollback()
            else:
                self.conn.commit()
        finally:
            self.conn.close()
        return False

    def execute(self, sql: str, params: Iterable[Any] = ()) -> DBResult:
        if IS_POSTGRES:
            sql = _qmark_to_percent(sql)
            cur = self.conn.cursor()
            cur.execute(sql, tuple(params))
            return DBResult(cur, True)
        return DBResult(self.conn.execute(sql, tuple(params)), False)

    def executescript(self, sql: str) -> None:
        if IS_POSTGRES:
            for statement in [s.strip() for s in sql.split(";") if s.strip()]:
                self.execute(statement)
        else:
            self.conn.executescript(sql)


def _qmark_to_percent(sql: str) -> str:
    # App SQL does not use literal '?' characters, so this simple conversion is safe here.
    return sql.replace("?", "%s")


def db() -> Database:
    return Database()


def scalar(conn: Database, sql: str, params: Iterable[Any] = ()) -> Any:
    row = conn.execute(sql, params).fetchone()
    if row is None:
        return None
    if isinstance(row, sqlite3.Row):
        return row[0]
    if isinstance(row, dict):
        return next(iter(row.values()))
    return row[0]


def insert_id(conn: Database, sql: str, params: Iterable[Any]) -> int:
    if IS_POSTGRES:
        row = conn.execute(sql.rstrip().rstrip(";") + " RETURNING id", params).fetchone()
        return int(row["id"])
    result = conn.execute(sql, params)
    return int(result.lastrowid)


def _sqlite_columns(conn: Database, table: str) -> set[str]:
    return {r["name"] for r in conn.execute(f"PRAGMA table_info({table})").fetchall()}


def _ensure_column(conn: Database, table: str, definition: str) -> None:
    name = definition.split()[0]
    if IS_POSTGRES:
        conn.execute(f"ALTER TABLE {table} ADD COLUMN IF NOT EXISTS {definition}")
    elif name not in _sqlite_columns(conn, table):
        conn.execute(f"ALTER TABLE {table} ADD COLUMN {definition}")


def init_db() -> None:
    validate_production_config()
    with db() as conn:
        # Serialize schema initialization across multiple production workers.
        if IS_POSTGRES:
            conn.execute("SELECT pg_advisory_xact_lock(?)", (824716534,))
        legacy_email_verification_migration = False
        if IS_POSTGRES:
            conn.executescript(
                """
                CREATE TABLE IF NOT EXISTS companies (
                    id BIGSERIAL PRIMARY KEY,
                    name TEXT NOT NULL,
                    company_type TEXT NOT NULL CHECK(company_type IN ('buyer','supplier','admin')),
                    location TEXT DEFAULT '', verified INTEGER NOT NULL DEFAULT 0,
                    active INTEGER NOT NULL DEFAULT 1, created_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS users (
                    id BIGSERIAL PRIMARY KEY,
                    company_id BIGINT NOT NULL REFERENCES companies(id) ON DELETE CASCADE,
                    name TEXT NOT NULL, email TEXT NOT NULL UNIQUE,
                    password_hash TEXT NOT NULL,
                    role TEXT NOT NULL CHECK(role IN ('buyer','supplier','admin')),
                    company_role TEXT NOT NULL DEFAULT 'owner',
                    email_verified INTEGER NOT NULL DEFAULT 0,
                    email_verified_at TEXT, session_version INTEGER NOT NULL DEFAULT 1,
                    last_login_at TEXT, active INTEGER NOT NULL DEFAULT 1, created_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS inventory (
                    id BIGSERIAL PRIMARY KEY,
                    company_id BIGINT NOT NULL REFERENCES companies(id) ON DELETE CASCADE,
                    brand TEXT NOT NULL DEFAULT '', part_number TEXT NOT NULL,
                    normalized_part TEXT NOT NULL, description TEXT NOT NULL DEFAULT '',
                    condition TEXT NOT NULL DEFAULT 'Not stated', quantity INTEGER NOT NULL DEFAULT 1,
                    location TEXT NOT NULL DEFAULT '', active INTEGER NOT NULL DEFAULT 1,
                    archived_at TEXT, deleted_at TEXT, created_at TEXT,
                    created_by_user_id BIGINT REFERENCES users(id) ON DELETE SET NULL,
                    updated_by_user_id BIGINT REFERENCES users(id) ON DELETE SET NULL,
                    last_confirmed_at TEXT, updated_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS rfqs (
                    id BIGSERIAL PRIMARY KEY, buyer_user_id BIGINT NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                    part_number TEXT NOT NULL, quantity INTEGER NOT NULL DEFAULT 1,
                    required_by TEXT DEFAULT '', delivery_location TEXT DEFAULT '', notes TEXT DEFAULT '',
                    status TEXT NOT NULL DEFAULT 'open', created_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS rfq_recipients (
                    id BIGSERIAL PRIMARY KEY, rfq_id BIGINT NOT NULL REFERENCES rfqs(id) ON DELETE CASCADE,
                    supplier_company_id BIGINT NOT NULL REFERENCES companies(id) ON DELETE CASCADE,
                    inventory_id BIGINT REFERENCES inventory(id) ON DELETE SET NULL, status TEXT NOT NULL DEFAULT 'sent',
                    quoted_price DOUBLE PRECISION, quoted_currency TEXT DEFAULT 'GBP', supplier_message TEXT DEFAULT '',
                    responded_at TEXT, accepted_at TEXT, UNIQUE(rfq_id, supplier_company_id)
                );
                CREATE TABLE IF NOT EXISTS wanted_requests (
                    id BIGSERIAL PRIMARY KEY, buyer_user_id BIGINT NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                    query TEXT NOT NULL, quantity INTEGER NOT NULL DEFAULT 1, notes TEXT DEFAULT '',
                    status TEXT NOT NULL DEFAULT 'open', created_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS wanted_responses (
                    id BIGSERIAL PRIMARY KEY, wanted_id BIGINT NOT NULL REFERENCES wanted_requests(id) ON DELETE CASCADE,
                    supplier_user_id BIGINT NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                    message TEXT NOT NULL, created_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS saved_searches (
                    id BIGSERIAL PRIMARY KEY, user_id BIGINT NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                    query TEXT NOT NULL, normalized_query TEXT NOT NULL, brand TEXT NOT NULL DEFAULT '',
                    condition TEXT NOT NULL DEFAULT '', active INTEGER NOT NULL DEFAULT 1,
                    last_alerted_at TEXT, created_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS saved_search_matches (
                    id BIGSERIAL PRIMARY KEY, saved_search_id BIGINT NOT NULL REFERENCES saved_searches(id) ON DELETE CASCADE,
                    inventory_id BIGINT NOT NULL REFERENCES inventory(id) ON DELETE CASCADE, notified_at TEXT NOT NULL,
                    UNIQUE(saved_search_id, inventory_id)
                );
                CREATE TABLE IF NOT EXISTS rfq_messages (
                    id BIGSERIAL PRIMARY KEY, recipient_id BIGINT NOT NULL REFERENCES rfq_recipients(id) ON DELETE CASCADE,
                    sender_user_id BIGINT NOT NULL REFERENCES users(id) ON DELETE CASCADE, body TEXT NOT NULL, created_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS wanted_matches (
                    id BIGSERIAL PRIMARY KEY, wanted_id BIGINT NOT NULL REFERENCES wanted_requests(id) ON DELETE CASCADE,
                    inventory_id BIGINT NOT NULL REFERENCES inventory(id) ON DELETE CASCADE, notified_at TEXT NOT NULL,
                    UNIQUE(wanted_id, inventory_id)
                );
                CREATE TABLE IF NOT EXISTS auth_tokens (
                    id BIGSERIAL PRIMARY KEY, token_hash TEXT NOT NULL UNIQUE,
                    purpose TEXT NOT NULL, user_id BIGINT REFERENCES users(id) ON DELETE CASCADE,
                    company_id BIGINT REFERENCES companies(id) ON DELETE CASCADE,
                    email TEXT NOT NULL DEFAULT '', company_role TEXT NOT NULL DEFAULT '',
                    expires_at TEXT NOT NULL, used_at TEXT, created_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS audit_logs (
                    id BIGSERIAL PRIMARY KEY, actor_user_id BIGINT REFERENCES users(id) ON DELETE SET NULL,
                    company_id BIGINT REFERENCES companies(id) ON DELETE SET NULL,
                    action TEXT NOT NULL, entity_type TEXT NOT NULL DEFAULT '', entity_id TEXT NOT NULL DEFAULT '',
                    metadata TEXT NOT NULL DEFAULT '{}', ip_address TEXT NOT NULL DEFAULT '', created_at TEXT NOT NULL
                );
                CREATE INDEX IF NOT EXISTS idx_inventory_normalized_part ON inventory(normalized_part);
                CREATE INDEX IF NOT EXISTS idx_inventory_company ON inventory(company_id);
                CREATE INDEX IF NOT EXISTS idx_auth_tokens_hash ON auth_tokens(token_hash);
                CREATE INDEX IF NOT EXISTS idx_audit_company_created ON audit_logs(company_id, created_at);
                CREATE INDEX IF NOT EXISTS idx_saved_search_user ON saved_searches(user_id, active);
                CREATE INDEX IF NOT EXISTS idx_rfq_messages_recipient ON rfq_messages(recipient_id, created_at);
                CREATE INDEX IF NOT EXISTS idx_wanted_matches_wanted ON wanted_matches(wanted_id);
                """
            )
        else:
            existing_user_columns = _sqlite_columns(conn, "users") if scalar(conn, "SELECT COUNT(*) FROM sqlite_master WHERE type='table' AND name='users'") else set()
            legacy_email_verification_migration = bool(existing_user_columns and "email_verified" not in existing_user_columns)
            conn.executescript(
                """
                CREATE TABLE IF NOT EXISTS companies (
                    id INTEGER PRIMARY KEY AUTOINCREMENT, name TEXT NOT NULL,
                    company_type TEXT NOT NULL CHECK(company_type IN ('buyer','supplier','admin')),
                    location TEXT DEFAULT '', verified INTEGER NOT NULL DEFAULT 0,
                    active INTEGER NOT NULL DEFAULT 1, created_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS users (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    company_id INTEGER NOT NULL REFERENCES companies(id) ON DELETE CASCADE,
                    name TEXT NOT NULL, email TEXT NOT NULL UNIQUE COLLATE NOCASE,
                    password_hash TEXT NOT NULL, role TEXT NOT NULL CHECK(role IN ('buyer','supplier','admin')),
                    company_role TEXT NOT NULL DEFAULT 'owner', email_verified INTEGER NOT NULL DEFAULT 0,
                    email_verified_at TEXT, session_version INTEGER NOT NULL DEFAULT 1, last_login_at TEXT,
                    active INTEGER NOT NULL DEFAULT 1, created_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS inventory (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    company_id INTEGER NOT NULL REFERENCES companies(id) ON DELETE CASCADE,
                    brand TEXT NOT NULL DEFAULT '', part_number TEXT NOT NULL, normalized_part TEXT NOT NULL,
                    description TEXT NOT NULL DEFAULT '', condition TEXT NOT NULL DEFAULT 'Not stated',
                    quantity INTEGER NOT NULL DEFAULT 1, location TEXT NOT NULL DEFAULT '', active INTEGER NOT NULL DEFAULT 1,
                    archived_at TEXT, deleted_at TEXT, created_at TEXT,
                    created_by_user_id INTEGER REFERENCES users(id) ON DELETE SET NULL,
                    updated_by_user_id INTEGER REFERENCES users(id) ON DELETE SET NULL,
                    last_confirmed_at TEXT, updated_at TEXT NOT NULL
                );
                CREATE INDEX IF NOT EXISTS idx_inventory_normalized_part ON inventory(normalized_part);
                CREATE INDEX IF NOT EXISTS idx_inventory_company ON inventory(company_id);
                CREATE TABLE IF NOT EXISTS rfqs (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    buyer_user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                    part_number TEXT NOT NULL, quantity INTEGER NOT NULL DEFAULT 1,
                    required_by TEXT DEFAULT '', delivery_location TEXT DEFAULT '', notes TEXT DEFAULT '',
                    status TEXT NOT NULL DEFAULT 'open', created_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS rfq_recipients (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    rfq_id INTEGER NOT NULL REFERENCES rfqs(id) ON DELETE CASCADE,
                    supplier_company_id INTEGER NOT NULL REFERENCES companies(id) ON DELETE CASCADE,
                    inventory_id INTEGER REFERENCES inventory(id) ON DELETE SET NULL, status TEXT NOT NULL DEFAULT 'sent',
                    quoted_price REAL, quoted_currency TEXT DEFAULT 'GBP', supplier_message TEXT DEFAULT '',
                    responded_at TEXT, accepted_at TEXT, UNIQUE(rfq_id, supplier_company_id)
                );
                CREATE TABLE IF NOT EXISTS wanted_requests (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    buyer_user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                    query TEXT NOT NULL, quantity INTEGER NOT NULL DEFAULT 1, notes TEXT DEFAULT '',
                    status TEXT NOT NULL DEFAULT 'open', created_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS wanted_responses (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    wanted_id INTEGER NOT NULL REFERENCES wanted_requests(id) ON DELETE CASCADE,
                    supplier_user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                    message TEXT NOT NULL, created_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS saved_searches (
                    id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE,
                    query TEXT NOT NULL, normalized_query TEXT NOT NULL, brand TEXT NOT NULL DEFAULT '',
                    condition TEXT NOT NULL DEFAULT '', active INTEGER NOT NULL DEFAULT 1,
                    last_alerted_at TEXT, created_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS saved_search_matches (
                    id INTEGER PRIMARY KEY AUTOINCREMENT, saved_search_id INTEGER NOT NULL REFERENCES saved_searches(id) ON DELETE CASCADE,
                    inventory_id INTEGER NOT NULL REFERENCES inventory(id) ON DELETE CASCADE, notified_at TEXT NOT NULL,
                    UNIQUE(saved_search_id, inventory_id)
                );
                CREATE TABLE IF NOT EXISTS rfq_messages (
                    id INTEGER PRIMARY KEY AUTOINCREMENT, recipient_id INTEGER NOT NULL REFERENCES rfq_recipients(id) ON DELETE CASCADE,
                    sender_user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE, body TEXT NOT NULL, created_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS wanted_matches (
                    id INTEGER PRIMARY KEY AUTOINCREMENT, wanted_id INTEGER NOT NULL REFERENCES wanted_requests(id) ON DELETE CASCADE,
                    inventory_id INTEGER NOT NULL REFERENCES inventory(id) ON DELETE CASCADE, notified_at TEXT NOT NULL,
                    UNIQUE(wanted_id, inventory_id)
                );
                CREATE TABLE IF NOT EXISTS auth_tokens (
                    id INTEGER PRIMARY KEY AUTOINCREMENT, token_hash TEXT NOT NULL UNIQUE,
                    purpose TEXT NOT NULL, user_id INTEGER REFERENCES users(id) ON DELETE CASCADE,
                    company_id INTEGER REFERENCES companies(id) ON DELETE CASCADE,
                    email TEXT NOT NULL DEFAULT '', company_role TEXT NOT NULL DEFAULT '',
                    expires_at TEXT NOT NULL, used_at TEXT, created_at TEXT NOT NULL
                );
                CREATE INDEX IF NOT EXISTS idx_auth_tokens_hash ON auth_tokens(token_hash);
                CREATE TABLE IF NOT EXISTS audit_logs (
                    id INTEGER PRIMARY KEY AUTOINCREMENT, actor_user_id INTEGER REFERENCES users(id) ON DELETE SET NULL,
                    company_id INTEGER REFERENCES companies(id) ON DELETE SET NULL,
                    action TEXT NOT NULL, entity_type TEXT NOT NULL DEFAULT '', entity_id TEXT NOT NULL DEFAULT '',
                    metadata TEXT NOT NULL DEFAULT '{}', ip_address TEXT NOT NULL DEFAULT '', created_at TEXT NOT NULL
                );
                CREATE INDEX IF NOT EXISTS idx_audit_company_created ON audit_logs(company_id, created_at);
                CREATE INDEX IF NOT EXISTS idx_saved_search_user ON saved_searches(user_id, active);
                CREATE INDEX IF NOT EXISTS idx_rfq_messages_recipient ON rfq_messages(recipient_id, created_at);
                CREATE INDEX IF NOT EXISTS idx_wanted_matches_wanted ON wanted_matches(wanted_id);
                """
            )
            # Safe additive migration from the original MVP schema.
            _ensure_column(conn, "users", "company_role TEXT NOT NULL DEFAULT 'owner'")
            _ensure_column(conn, "users", "email_verified INTEGER NOT NULL DEFAULT 0")
            _ensure_column(conn, "users", "email_verified_at TEXT")
            _ensure_column(conn, "users", "session_version INTEGER NOT NULL DEFAULT 1")
            _ensure_column(conn, "users", "last_login_at TEXT")
            _ensure_column(conn, "inventory", "archived_at TEXT")
            _ensure_column(conn, "inventory", "deleted_at TEXT")
            _ensure_column(conn, "inventory", "created_at TEXT")
            _ensure_column(conn, "inventory", "created_by_user_id INTEGER")
            _ensure_column(conn, "inventory", "updated_by_user_id INTEGER")

        # Phase 2 additive migrations for both SQLite and PostgreSQL deployments.
        _ensure_column(conn, "inventory", "last_confirmed_at TEXT")
        _ensure_column(conn, "rfq_recipients", "accepted_at TEXT")

        # Phase 3 supplier-ingestion metadata and automation tables.
        if IS_POSTGRES:
            conn.executescript(
                """
                CREATE TABLE IF NOT EXISTS inventory_import_profiles (
                    id BIGSERIAL PRIMARY KEY, company_id BIGINT NOT NULL REFERENCES companies(id) ON DELETE CASCADE,
                    name TEXT NOT NULL, mapping_json TEXT NOT NULL DEFAULT '{}', created_by_user_id BIGINT REFERENCES users(id) ON DELETE SET NULL,
                    created_at TEXT NOT NULL, updated_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS inventory_feeds (
                    id BIGSERIAL PRIMARY KEY, company_id BIGINT NOT NULL REFERENCES companies(id) ON DELETE CASCADE,
                    name TEXT NOT NULL, feed_type TEXT NOT NULL, enabled INTEGER NOT NULL DEFAULT 1,
                    schedule_minutes INTEGER NOT NULL DEFAULT 360, next_run_at TEXT, last_run_at TEXT, last_status TEXT NOT NULL DEFAULT '',
                    config_json TEXT NOT NULL DEFAULT '{}', secret_blob TEXT NOT NULL DEFAULT '', mapping_json TEXT NOT NULL DEFAULT '{}',
                    import_mode TEXT NOT NULL DEFAULT 'sync', duplicate_policy TEXT NOT NULL DEFAULT 'update',
                    ingest_token_hash TEXT NOT NULL DEFAULT '', ingest_token_blob TEXT NOT NULL DEFAULT '',
                    created_by_user_id BIGINT REFERENCES users(id) ON DELETE SET NULL, locked_at TEXT, created_at TEXT NOT NULL, updated_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS inventory_import_jobs (
                    id BIGSERIAL PRIMARY KEY, company_id BIGINT NOT NULL REFERENCES companies(id) ON DELETE CASCADE,
                    user_id BIGINT REFERENCES users(id) ON DELETE SET NULL, feed_id BIGINT REFERENCES inventory_feeds(id) ON DELETE SET NULL,
                    source_type TEXT NOT NULL, source_name TEXT NOT NULL DEFAULT '', import_mode TEXT NOT NULL DEFAULT 'upsert',
                    duplicate_policy TEXT NOT NULL DEFAULT 'update', status TEXT NOT NULL DEFAULT 'running', total_rows INTEGER NOT NULL DEFAULT 0,
                    inserted_rows INTEGER NOT NULL DEFAULT 0, updated_rows INTEGER NOT NULL DEFAULT 0, archived_rows INTEGER NOT NULL DEFAULT 0,
                    duplicate_rows INTEGER NOT NULL DEFAULT 0, skipped_rows INTEGER NOT NULL DEFAULT 0, error_rows INTEGER NOT NULL DEFAULT 0,
                    report_json TEXT NOT NULL DEFAULT '{}', error_summary TEXT NOT NULL DEFAULT '', started_at TEXT NOT NULL, completed_at TEXT
                );
                CREATE TABLE IF NOT EXISTS inventory_import_errors (
                    id BIGSERIAL PRIMARY KEY, job_id BIGINT NOT NULL REFERENCES inventory_import_jobs(id) ON DELETE CASCADE,
                    row_number INTEGER, severity TEXT NOT NULL DEFAULT 'error', code TEXT NOT NULL DEFAULT '', message TEXT NOT NULL,
                    raw_json TEXT NOT NULL DEFAULT '{}', created_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS inventory_import_staging (
                    id BIGSERIAL PRIMARY KEY, company_id BIGINT NOT NULL REFERENCES companies(id) ON DELETE CASCADE,
                    user_id BIGINT NOT NULL REFERENCES users(id) ON DELETE CASCADE, filename TEXT NOT NULL, file_data BYTEA NOT NULL,
                    headers_json TEXT NOT NULL DEFAULT '[]', expires_at TEXT NOT NULL, created_at TEXT NOT NULL
                );
                CREATE INDEX IF NOT EXISTS idx_inventory_feeds_due ON inventory_feeds(enabled, next_run_at);
                CREATE INDEX IF NOT EXISTS idx_import_jobs_company_started ON inventory_import_jobs(company_id, started_at);
                CREATE INDEX IF NOT EXISTS idx_import_errors_job ON inventory_import_errors(job_id);
                """
            )
        else:
            conn.executescript(
                """
                CREATE TABLE IF NOT EXISTS inventory_import_profiles (
                    id INTEGER PRIMARY KEY AUTOINCREMENT, company_id INTEGER NOT NULL REFERENCES companies(id) ON DELETE CASCADE,
                    name TEXT NOT NULL, mapping_json TEXT NOT NULL DEFAULT '{}', created_by_user_id INTEGER REFERENCES users(id) ON DELETE SET NULL,
                    created_at TEXT NOT NULL, updated_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS inventory_feeds (
                    id INTEGER PRIMARY KEY AUTOINCREMENT, company_id INTEGER NOT NULL REFERENCES companies(id) ON DELETE CASCADE,
                    name TEXT NOT NULL, feed_type TEXT NOT NULL, enabled INTEGER NOT NULL DEFAULT 1,
                    schedule_minutes INTEGER NOT NULL DEFAULT 360, next_run_at TEXT, last_run_at TEXT, last_status TEXT NOT NULL DEFAULT '',
                    config_json TEXT NOT NULL DEFAULT '{}', secret_blob TEXT NOT NULL DEFAULT '', mapping_json TEXT NOT NULL DEFAULT '{}',
                    import_mode TEXT NOT NULL DEFAULT 'sync', duplicate_policy TEXT NOT NULL DEFAULT 'update',
                    ingest_token_hash TEXT NOT NULL DEFAULT '', ingest_token_blob TEXT NOT NULL DEFAULT '',
                    created_by_user_id INTEGER REFERENCES users(id) ON DELETE SET NULL, locked_at TEXT, created_at TEXT NOT NULL, updated_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS inventory_import_jobs (
                    id INTEGER PRIMARY KEY AUTOINCREMENT, company_id INTEGER NOT NULL REFERENCES companies(id) ON DELETE CASCADE,
                    user_id INTEGER REFERENCES users(id) ON DELETE SET NULL, feed_id INTEGER REFERENCES inventory_feeds(id) ON DELETE SET NULL,
                    source_type TEXT NOT NULL, source_name TEXT NOT NULL DEFAULT '', import_mode TEXT NOT NULL DEFAULT 'upsert',
                    duplicate_policy TEXT NOT NULL DEFAULT 'update', status TEXT NOT NULL DEFAULT 'running', total_rows INTEGER NOT NULL DEFAULT 0,
                    inserted_rows INTEGER NOT NULL DEFAULT 0, updated_rows INTEGER NOT NULL DEFAULT 0, archived_rows INTEGER NOT NULL DEFAULT 0,
                    duplicate_rows INTEGER NOT NULL DEFAULT 0, skipped_rows INTEGER NOT NULL DEFAULT 0, error_rows INTEGER NOT NULL DEFAULT 0,
                    report_json TEXT NOT NULL DEFAULT '{}', error_summary TEXT NOT NULL DEFAULT '', started_at TEXT NOT NULL, completed_at TEXT
                );
                CREATE TABLE IF NOT EXISTS inventory_import_errors (
                    id INTEGER PRIMARY KEY AUTOINCREMENT, job_id INTEGER NOT NULL REFERENCES inventory_import_jobs(id) ON DELETE CASCADE,
                    row_number INTEGER, severity TEXT NOT NULL DEFAULT 'error', code TEXT NOT NULL DEFAULT '', message TEXT NOT NULL,
                    raw_json TEXT NOT NULL DEFAULT '{}', created_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS inventory_import_staging (
                    id INTEGER PRIMARY KEY AUTOINCREMENT, company_id INTEGER NOT NULL REFERENCES companies(id) ON DELETE CASCADE,
                    user_id INTEGER NOT NULL REFERENCES users(id) ON DELETE CASCADE, filename TEXT NOT NULL, file_data BLOB NOT NULL,
                    headers_json TEXT NOT NULL DEFAULT '[]', expires_at TEXT NOT NULL, created_at TEXT NOT NULL
                );
                CREATE INDEX IF NOT EXISTS idx_inventory_feeds_due ON inventory_feeds(enabled, next_run_at);
                CREATE INDEX IF NOT EXISTS idx_import_jobs_company_started ON inventory_import_jobs(company_id, started_at);
                CREATE INDEX IF NOT EXISTS idx_import_errors_job ON inventory_import_errors(job_id);
                """
            )
        _ensure_column(conn, "inventory", "source_feed_id INTEGER")
        _ensure_column(conn, "inventory", "source_key TEXT")
        _ensure_column(conn, "inventory", "last_import_job_id INTEGER")

        # Phase 4: canonical catalogue, technical relationships, photo ID and sourcing intelligence.
        if IS_POSTGRES:
            conn.executescript(
                """
                CREATE TABLE IF NOT EXISTS catalog_parts (
                    id BIGSERIAL PRIMARY KEY, manufacturer TEXT NOT NULL DEFAULT '', part_number TEXT NOT NULL,
                    normalized_part TEXT NOT NULL, name TEXT NOT NULL DEFAULT '', product_family TEXT NOT NULL DEFAULT '',
                    description TEXT NOT NULL DEFAULT '', lifecycle_status TEXT NOT NULL DEFAULT 'unknown',
                    notes TEXT NOT NULL DEFAULT '', source TEXT NOT NULL DEFAULT '', verified INTEGER NOT NULL DEFAULT 0,
                    created_by_user_id BIGINT REFERENCES users(id) ON DELETE SET NULL,
                    created_at TEXT NOT NULL, updated_at TEXT NOT NULL, UNIQUE(manufacturer, normalized_part)
                );
                CREATE TABLE IF NOT EXISTS catalog_aliases (
                    id BIGSERIAL PRIMARY KEY, part_id BIGINT NOT NULL REFERENCES catalog_parts(id) ON DELETE CASCADE,
                    alias TEXT NOT NULL, normalized_alias TEXT NOT NULL, source TEXT NOT NULL DEFAULT '',
                    verified INTEGER NOT NULL DEFAULT 0, created_at TEXT NOT NULL, UNIQUE(part_id, normalized_alias)
                );
                CREATE TABLE IF NOT EXISTS catalog_relations (
                    id BIGSERIAL PRIMARY KEY, from_part_id BIGINT NOT NULL REFERENCES catalog_parts(id) ON DELETE CASCADE,
                    to_part_id BIGINT NOT NULL REFERENCES catalog_parts(id) ON DELETE CASCADE, relation_type TEXT NOT NULL,
                    confidence DOUBLE PRECISION NOT NULL DEFAULT 1.0, notes TEXT NOT NULL DEFAULT '', source TEXT NOT NULL DEFAULT '',
                    verified INTEGER NOT NULL DEFAULT 0, created_by_user_id BIGINT REFERENCES users(id) ON DELETE SET NULL,
                    created_at TEXT NOT NULL, updated_at TEXT NOT NULL, UNIQUE(from_part_id, to_part_id, relation_type)
                );
                CREATE TABLE IF NOT EXISTS photo_identifications (
                    id BIGSERIAL PRIMARY KEY, user_id BIGINT REFERENCES users(id) ON DELETE SET NULL,
                    company_id BIGINT REFERENCES companies(id) ON DELETE SET NULL, provider TEXT NOT NULL DEFAULT '',
                    original_filename TEXT NOT NULL DEFAULT '', visible_text TEXT NOT NULL DEFAULT '',
                    identified_manufacturer TEXT NOT NULL DEFAULT '', identified_part_number TEXT NOT NULL DEFAULT '',
                    confidence DOUBLE PRECISION, matched_part_id BIGINT REFERENCES catalog_parts(id) ON DELETE SET NULL,
                    result_json TEXT NOT NULL DEFAULT '{}', created_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS sourcing_queries (
                    id BIGSERIAL PRIMARY KEY, user_id BIGINT REFERENCES users(id) ON DELETE SET NULL,
                    company_id BIGINT REFERENCES companies(id) ON DELETE SET NULL, query TEXT NOT NULL,
                    result_json TEXT NOT NULL DEFAULT '{}', created_at TEXT NOT NULL
                );
                CREATE INDEX IF NOT EXISTS idx_catalog_normalized_part ON catalog_parts(normalized_part);
                CREATE INDEX IF NOT EXISTS idx_catalog_alias_normalized ON catalog_aliases(normalized_alias);
                CREATE INDEX IF NOT EXISTS idx_catalog_relations_from ON catalog_relations(from_part_id);
                CREATE INDEX IF NOT EXISTS idx_catalog_relations_to ON catalog_relations(to_part_id);
                """
            )
        else:
            conn.executescript(
                """
                CREATE TABLE IF NOT EXISTS catalog_parts (
                    id INTEGER PRIMARY KEY AUTOINCREMENT, manufacturer TEXT NOT NULL DEFAULT '', part_number TEXT NOT NULL,
                    normalized_part TEXT NOT NULL, name TEXT NOT NULL DEFAULT '', product_family TEXT NOT NULL DEFAULT '',
                    description TEXT NOT NULL DEFAULT '', lifecycle_status TEXT NOT NULL DEFAULT 'unknown',
                    notes TEXT NOT NULL DEFAULT '', source TEXT NOT NULL DEFAULT '', verified INTEGER NOT NULL DEFAULT 0,
                    created_by_user_id INTEGER REFERENCES users(id) ON DELETE SET NULL,
                    created_at TEXT NOT NULL, updated_at TEXT NOT NULL, UNIQUE(manufacturer, normalized_part)
                );
                CREATE TABLE IF NOT EXISTS catalog_aliases (
                    id INTEGER PRIMARY KEY AUTOINCREMENT, part_id INTEGER NOT NULL REFERENCES catalog_parts(id) ON DELETE CASCADE,
                    alias TEXT NOT NULL, normalized_alias TEXT NOT NULL, source TEXT NOT NULL DEFAULT '',
                    verified INTEGER NOT NULL DEFAULT 0, created_at TEXT NOT NULL, UNIQUE(part_id, normalized_alias)
                );
                CREATE TABLE IF NOT EXISTS catalog_relations (
                    id INTEGER PRIMARY KEY AUTOINCREMENT, from_part_id INTEGER NOT NULL REFERENCES catalog_parts(id) ON DELETE CASCADE,
                    to_part_id INTEGER NOT NULL REFERENCES catalog_parts(id) ON DELETE CASCADE, relation_type TEXT NOT NULL,
                    confidence REAL NOT NULL DEFAULT 1.0, notes TEXT NOT NULL DEFAULT '', source TEXT NOT NULL DEFAULT '',
                    verified INTEGER NOT NULL DEFAULT 0, created_by_user_id INTEGER REFERENCES users(id) ON DELETE SET NULL,
                    created_at TEXT NOT NULL, updated_at TEXT NOT NULL, UNIQUE(from_part_id, to_part_id, relation_type)
                );
                CREATE TABLE IF NOT EXISTS photo_identifications (
                    id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER REFERENCES users(id) ON DELETE SET NULL,
                    company_id INTEGER REFERENCES companies(id) ON DELETE SET NULL, provider TEXT NOT NULL DEFAULT '',
                    original_filename TEXT NOT NULL DEFAULT '', visible_text TEXT NOT NULL DEFAULT '',
                    identified_manufacturer TEXT NOT NULL DEFAULT '', identified_part_number TEXT NOT NULL DEFAULT '',
                    confidence REAL, matched_part_id INTEGER REFERENCES catalog_parts(id) ON DELETE SET NULL,
                    result_json TEXT NOT NULL DEFAULT '{}', created_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS sourcing_queries (
                    id INTEGER PRIMARY KEY AUTOINCREMENT, user_id INTEGER REFERENCES users(id) ON DELETE SET NULL,
                    company_id INTEGER REFERENCES companies(id) ON DELETE SET NULL, query TEXT NOT NULL,
                    result_json TEXT NOT NULL DEFAULT '{}', created_at TEXT NOT NULL
                );
                CREATE INDEX IF NOT EXISTS idx_catalog_normalized_part ON catalog_parts(normalized_part);
                CREATE INDEX IF NOT EXISTS idx_catalog_alias_normalized ON catalog_aliases(normalized_alias);
                CREATE INDEX IF NOT EXISTS idx_catalog_relations_from ON catalog_relations(from_part_id);
                CREATE INDEX IF NOT EXISTS idx_catalog_relations_to ON catalog_relations(to_part_id);
                """
            )
        _ensure_column(conn, "inventory", "canonical_part_id INTEGER")
        _ensure_column(conn, "inventory", "catalog_match_method TEXT")
        _ensure_column(conn, "inventory", "catalog_match_confidence REAL")

        # Treat accounts from the earlier SQLite MVP as verified owners exactly once during schema migration.
        conn.execute("UPDATE users SET company_role='owner' WHERE company_role IS NULL OR company_role='' ")
        if legacy_email_verification_migration:
            conn.execute("UPDATE users SET email_verified=1, email_verified_at=COALESCE(email_verified_at, created_at) WHERE email_verified=0")
        conn.execute("UPDATE inventory SET created_at=COALESCE(created_at, updated_at) WHERE created_at IS NULL")
        conn.execute("UPDATE inventory SET last_confirmed_at=COALESCE(last_confirmed_at, updated_at, created_at) WHERE last_confirmed_at IS NULL")

        admin_email = os.getenv("SEED_ADMIN_EMAIL", "admin@controlsexchange.local").strip().lower()
        admin_password = os.getenv("SEED_ADMIN_PASSWORD", "ChangeMe123!")
        seed_admin = os.getenv("SEED_ADMIN", "true" if ENVIRONMENT != "production" else "false").lower() == "true"
        if seed_admin and not scalar(conn, "SELECT COUNT(*) FROM users WHERE role='admin'"):
            company_id = insert_id(conn,
                "INSERT INTO companies(name,company_type,location,verified,active,created_at) VALUES(?,?,?,?,?,?)",
                ("Controls Exchange", "admin", "United Kingdom", 1, 1, now_iso()))
            conn.execute(
                "INSERT INTO users(company_id,name,email,password_hash,role,company_role,email_verified,email_verified_at,active,created_at) VALUES(?,?,?,?,?,?,?,?,?,?)",
                (company_id, "Platform Admin", admin_email, hash_password(admin_password), "admin", "owner", 1, now_iso(), 1, now_iso()),
            )

        seed_demo = os.getenv("SEED_DEMO_DATA", "true" if ENVIRONMENT != "production" else "false").lower() == "true"
        if seed_demo and not scalar(conn, "SELECT COUNT(*) FROM companies WHERE name='Demo Controls Ltd'"):
            supplier_company = insert_id(conn,
                "INSERT INTO companies(name,company_type,location,verified,active,created_at) VALUES(?,?,?,?,?,?)",
                ("Demo Controls Ltd", "supplier", "Birmingham, UK", 1, 1, now_iso()))
            supplier_user = insert_id(conn,
                "INSERT INTO users(company_id,name,email,password_hash,role,company_role,email_verified,email_verified_at,active,created_at) VALUES(?,?,?,?,?,?,?,?,?,?)",
                (supplier_company, "Demo Supplier", "supplier@example.com", hash_password("Supplier123!"), "supplier", "owner", 1, now_iso(), 1, now_iso()))
            demo_items = [
                ("Trend", "IQ233/UNB/230VAC", "IQ2 series universal network controller", "New old stock", 3, "Birmingham, UK"),
                ("Honeywell", "XFL521B", "Excel 500 I/O module", "New old stock", 6, "Birmingham, UK"),
                ("Johnson Controls", "MS-NAE5510-2", "Metasys network automation engine", "Used / tested", 2, "Birmingham, UK"),
                ("Schneider", "TAC XENTA 401", "Programmable HVAC controller", "Refurbished", 5, "Birmingham, UK"),
                ("Siemens", "PXC64-U", "DESIGO automation station", "Refurbished", 2, "Birmingham, UK"),
                ("ABB / Cylon", "UC32.24", "Unitron UC32 field controller", "Used / tested", 7, "Birmingham, UK"),
            ]
            for brand, part, desc, condition, qty, loc in demo_items:
                conn.execute(
                    "INSERT INTO inventory(company_id,brand,part_number,normalized_part,description,condition,quantity,location,active,created_at,created_by_user_id,updated_by_user_id,last_confirmed_at,updated_at) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                    (supplier_company, brand, part, normalize(part), desc, condition, qty, loc, 1, now_iso(), supplier_user, supplier_user, now_iso(), now_iso()),
                )
        if seed_demo and not scalar(conn, "SELECT COUNT(*) FROM users WHERE email='buyer@example.com'"):
            buyer_company = insert_id(conn,
                "INSERT INTO companies(name,company_type,location,verified,active,created_at) VALUES(?,?,?,?,?,?)",
                ("Demo Integrator Ltd", "buyer", "London, UK", 1, 1, now_iso()))
            conn.execute(
                "INSERT INTO users(company_id,name,email,password_hash,role,company_role,email_verified,email_verified_at,active,created_at) VALUES(?,?,?,?,?,?,?,?,?,?)",
                (buyer_company, "Demo Buyer", "buyer@example.com", hash_password("Buyer123!"), "buyer", "owner", 1, now_iso(), 1, now_iso()),
            )


def can(user: Any, permission: str) -> bool:
    if not user:
        return False
    if user["role"] == "admin":
        return True
    return permission in ROLE_PERMISSIONS.get(user["company_role"] or "viewer", set())


def issue_token(conn: Database, purpose: str, *, user_id: int | None = None, company_id: int | None = None,
                email: str = "", company_role: str = "", hours: int = 24) -> str:
    raw = secrets.token_urlsafe(32)
    expires = (now_dt() + timedelta(hours=hours)).isoformat(timespec="seconds")
    conn.execute(
        "INSERT INTO auth_tokens(token_hash,purpose,user_id,company_id,email,company_role,expires_at,created_at) VALUES(?,?,?,?,?,?,?,?)",
        (token_hash(raw), purpose, user_id, company_id, email.lower().strip(), company_role, expires, now_iso()),
    )
    return raw


def get_valid_token(conn: Database, raw: str, purpose: str):
    row = conn.execute(
        "SELECT * FROM auth_tokens WHERE token_hash=? AND purpose=? AND used_at IS NULL",
        (token_hash(raw), purpose),
    ).fetchone()
    if not row:
        return None
    try:
        if datetime.fromisoformat(row["expires_at"]) < now_dt():
            return None
    except Exception:
        return None
    return row


def consume_token(conn: Database, token_id: int) -> None:
    conn.execute("UPDATE auth_tokens SET used_at=? WHERE id=?", (now_iso(), token_id))


def request_ip(request: Any) -> str:
    forwarded = request.headers.get("x-forwarded-for", "") if request else ""
    if forwarded:
        return forwarded.split(",")[0].strip()[:64]
    return (request.client.host if request and request.client else "")[:64]


def audit(conn: Database, request: Any, actor: Any, action: str, entity_type: str = "", entity_id: Any = "", metadata: Optional[dict] = None, company_id: int | None = None) -> None:
    if company_id is None and actor:
        company_id = actor["company_id"]
    actor_id = actor["id"] if actor else None
    conn.execute(
        "INSERT INTO audit_logs(actor_user_id,company_id,action,entity_type,entity_id,metadata,ip_address,created_at) VALUES(?,?,?,?,?,?,?,?)",
        (actor_id, company_id, action, entity_type, str(entity_id or ""), json.dumps(metadata or {}, separators=(",", ":")), request_ip(request), now_iso()),
    )


def send_email(to_email: str, subject: str, body: str, html: str | None = None) -> bool:
    provider = os.getenv("EMAIL_PROVIDER", "").strip().lower()
    if not provider:
        provider = "smtp" if os.getenv("SMTP_HOST", "").strip() else "log"
    sender = os.getenv("EMAIL_FROM", os.getenv("SMTP_FROM", "hello@controlsexchange.example")).strip()
    try:
        if provider == "resend":
            api_key = os.getenv("RESEND_API_KEY", "").strip()
            if not api_key:
                raise RuntimeError("RESEND_API_KEY is not configured")
            payload = {"from": sender, "to": [to_email], "subject": subject, "text": body}
            if html:
                payload["html"] = html
            response = httpx.post(
                "https://api.resend.com/emails", json=payload,
                headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"}, timeout=15,
            )
            response.raise_for_status()
            return True
        if provider == "smtp":
            host = os.getenv("SMTP_HOST", "").strip()
            if not host:
                raise RuntimeError("SMTP_HOST is not configured")
            port = int(os.getenv("SMTP_PORT", "587"))
            username = os.getenv("SMTP_USER", "").strip()
            password = os.getenv("SMTP_PASSWORD", "")
            use_tls = os.getenv("SMTP_TLS", "true").lower() == "true"
            msg = EmailMessage()
            msg["From"] = sender
            msg["To"] = to_email
            msg["Subject"] = subject
            msg.set_content(body)
            if html:
                msg.add_alternative(html, subtype="html")
            if port == 465:
                with smtplib.SMTP_SSL(host, port, context=ssl.create_default_context(), timeout=15) as smtp:
                    if username:
                        smtp.login(username, password)
                    smtp.send_message(msg)
            else:
                with smtplib.SMTP(host, port, timeout=15) as smtp:
                    if use_tls:
                        smtp.starttls(context=ssl.create_default_context())
                    if username:
                        smtp.login(username, password)
                    smtp.send_message(msg)
            return True
        outbox = DATA_DIR / "email_outbox.log"
        with outbox.open("a", encoding="utf-8") as f:
            f.write("\n" + "=" * 72 + "\n")
            f.write(f"TO: {to_email}\nSUBJECT: {subject}\n\n{body}\n")
        return True
    except Exception:
        logger.exception("email_send_failed", extra={"to": to_email, "subject": subject})
        return False


def parse_inventory_file(filename: str, raw: bytes) -> list[dict]:
    ext = Path(filename or "").suffix.lower()
    rows: list[dict] = []
    if ext == ".csv":
        text = raw.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        rows = [dict(r) for r in csv.DictReader(io.StringIO(text), dialect=dialect)]
    elif ext in {".xlsx", ".xlsm"}:
        wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        ws = wb.active
        iterator = ws.iter_rows(values_only=True)
        try:
            headers = [str(v).strip() if v is not None else "" for v in next(iterator)]
        except StopIteration:
            return []
        for values in iterator:
            rows.append({headers[i]: values[i] if i < len(values) else None for i in range(len(headers))})
    elif ext == ".xls":
        raise ValueError("Legacy .xls is not supported. Save/export it as .xlsx or .csv.")
    else:
        raise ValueError("Please upload a CSV or XLSX file.")

    aliases = {
        "brand": {"brand", "manufacturer", "make", "mfr", "vendor"},
        "part_number": {"partnumber", "partno", "part", "sku", "productcode", "stockcode", "model"},
        "description": {"description", "productdescription", "name", "productname", "details"},
        "condition": {"condition", "state", "stockcondition"},
        "quantity": {"quantity", "qty", "stock", "onhand", "available", "stockqty"},
        "location": {"location", "warehouse", "site", "city", "stocklocation"},
    }
    key_norm = lambda v: re.sub(r"[^a-z0-9]", "", (v or "").lower())
    if not rows:
        return []
    header_map: dict[str, str] = {}
    for original in rows[0].keys():
        nk = key_norm(str(original))
        for canonical, opts in aliases.items():
            if nk in opts and canonical not in header_map:
                header_map[canonical] = original
    if "part_number" not in header_map:
        raise ValueError("Could not find a part-number column. Use Part Number, Part No, SKU, Stock Code or Model.")
    cleaned = []
    for row in rows:
        part = str(row.get(header_map["part_number"], "") or "").strip()
        if not part:
            continue
        qraw = row.get(header_map.get("quantity", ""), 1)
        try:
            qty = max(0, int(float(str(qraw or 1).replace(",", ""))))
        except Exception:
            qty = 1
        cleaned.append({
            "brand": str(row.get(header_map.get("brand", ""), "") or "").strip(),
            "part_number": part,
            "description": str(row.get(header_map.get("description", ""), "") or "").strip(),
            "condition": str(row.get(header_map.get("condition", ""), "") or "Not stated").strip() or "Not stated",
            "quantity": qty,
            "location": str(row.get(header_map.get("location", ""), "") or "").strip(),
        })
    return cleaned
